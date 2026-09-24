#!/usr/bin/env python3
"""
qatarcid_scraper.py  (Scrapling + curl_cffi edition)

Flow:
  1. Scrapling StealthySession (solve_cloudflare=True, persistent profile) loads
     a page once, clears the Cloudflare Turnstile interstitial, and we pull
     cookies + UA + the theme_scriptspf nonce out of the live browser context.
  2. All bulk traffic uses curl_cffi (impersonate="chrome") so the TLS/JA3
     fingerprint matches a real Chrome -- cf_clearance can be bound to it.
  3. Listing sweep (AJAX, pfget_listitems) -> detail sweep -> Excel.
     Checkpointed; resumable.

Install:
    pip install "scrapling[fetchers]" curl_cffi openpyxl
    scrapling install

Usage:
    python qatarcid_scraper.py --probe                  # nonce, page size test, wp-json probe
    python qatarcid_scraper.py --page-size 75           # after probe confirms 75 works
    python qatarcid_scraper.py --workers 3              # parallel detail fetches
    python qatarcid_scraper.py --headed                 # if headless CF solve fails
    python qatarcid_scraper.py --export-only
"""

import argparse
import json
import math
import os
import re
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from html import unescape

from curl_cffi import requests as cr

BASE = "https://www.qatarcid.com"
AJAX_URL = f"{BASE}/wp-content/plugins/pointfindercoreelements/includes/pfajaxhandler.php"
CHECKPOINT_PATH = "qatarcid_checkpoint.json"
OUTPUT_XLSX = "qatarcid_companies.xlsx"
PROFILE_DIR = os.path.abspath("./scrapling_profile")

# A real listing permalink carries theme_scriptspf; homepage may not.
NONCE_URLS = [
    f"{BASE}/listing/elite-business-solution-service-2/",
    f"{BASE}/",
]

REQUEST_TIMEOUT = 30
RETRY_LIMIT = 3
RETRY_BACKOFF_SECONDS = 5
POLITE_DELAY_SECONDS = 1.5
DEFAULT_RETRY_AFTER_SECONDS = 30

LISTING_HREF_RE = re.compile(
    r'<li class="pflist-itemtitle[^"]*"><a href="(https://www\.qatarcid\.com/listing/[^"]+/)">([^<]+)</a></li>'
)
FOUNDPOSTS_RE = re.compile(r'data-foundposts="(\d+)"')
NONCE_RE = re.compile(r'var\s+theme_scriptspf\s*=\s*(\{.*?\});', re.DOTALL)
DETAIL_FIELD_RE = re.compile(
    r'<span class="pf-ftitle">([^<]+?)\s*:?\s*</span>\s*<span class="pfdetail-ftext">(.*?)</span>',
    re.DOTALL,
)
TAG_STRIP_RE = re.compile(r"<[^>]+>")


CFEMAIL_RE = re.compile(r'<(\w+)[^>]*data-cfemail="([0-9a-fA-F]+)"[^>]*>.*?</\1>', re.DOTALL)
CFEMAIL_HREF_RE = re.compile(r'<a[^>]*/cdn-cgi/l/email-protection#([0-9a-fA-F]+)[^>]*>.*?</a>', re.DOTALL)


def decode_cfemail(hexstr):
    """Cloudflare email obfuscation: first byte is an XOR key for the rest."""
    try:
        key = int(hexstr[:2], 16)
        return "".join(chr(int(hexstr[i:i + 2], 16) ^ key) for i in range(2, len(hexstr), 2))
    except ValueError:
        return ""


def decode_cf_emails(html):
    html = CFEMAIL_RE.sub(lambda m: decode_cfemail(m.group(2)), html)
    return CFEMAIL_HREF_RE.sub(lambda m: decode_cfemail(m.group(1)), html)


def strip_tags(s):
    return unescape(TAG_STRIP_RE.sub("", s)).strip()


def parse_listing_cards(html):
    return [(u, unescape(n).strip()) for u, n in LISTING_HREF_RE.findall(html)]


def parse_found_posts(html):
    m = FOUNDPOSTS_RE.search(html)
    return int(m.group(1)) if m else None


def parse_detail_fields(html):
    html = decode_cf_emails(html)
    out = {}
    for label, raw in DETAIL_FIELD_RE.findall(html):
        label = strip_tags(label)
        if label:
            out[label] = strip_tags(raw)
    return out


# ---------------------------------------------------------------------------
# Checkpoint
# ---------------------------------------------------------------------------
@dataclass
class ScrapeState:
    total_found: int = None
    total_pages: int = None
    page_size: int = 14
    pages_done: set = field(default_factory=set)
    listing_urls: dict = field(default_factory=dict)
    detail_done: set = field(default_factory=set)
    records: dict = field(default_factory=dict)

    @classmethod
    def load(cls, path):
        if not os.path.isfile(path):
            return cls()
        with open(path, "r", encoding="utf-8") as f:
            raw = json.load(f)
        return cls(
            total_found=raw.get("total_found"),
            total_pages=raw.get("total_pages"),
            page_size=raw.get("page_size", 14),
            pages_done=set(raw.get("pages_done", [])),
            listing_urls=raw.get("listing_urls", {}),
            detail_done=set(raw.get("detail_done", [])),
            records=raw.get("records", {}),
        )

    def save(self, path):
        tmp = path + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump({
                "total_found": self.total_found,
                "total_pages": self.total_pages,
                "page_size": self.page_size,
                "pages_done": sorted(self.pages_done),
                "listing_urls": self.listing_urls,
                "detail_done": sorted(self.detail_done),
                "records": self.records,
            }, f, ensure_ascii=False, indent=1)
        os.replace(tmp, path)


# ---------------------------------------------------------------------------
# Cloudflare bootstrap via Scrapling
# ---------------------------------------------------------------------------
def solve_with_scrapling(headless=True):
    """Returns (cookies_dict, user_agent, html). Uses page_action to read the
    live Playwright context so we don't depend on Scrapling's Response API."""
    from scrapling.fetchers import StealthySession

    captured = {}

    def grab(page):
        # Wait for the challenge to clear, then snapshot state.
        deadline = time.time() + 60
        while time.time() < deadline:
            names = {c["name"] for c in page.context.cookies()}
            html = page.content()
            if "cf_clearance" in names and "Performing security verification" not in html:
                break
            page.wait_for_timeout(1500)
        captured["cookies"] = {c["name"]: c["value"] for c in page.context.cookies()}
        captured["ua"] = page.evaluate("navigator.userAgent")
        captured["html"] = page.content()
        return page

    html = ""
    with StealthySession(
        headless=headless,
        solve_cloudflare=True,
        user_data_dir=PROFILE_DIR,      # persist clearance between runs
        network_idle=True,
    ) as session:
        for url in NONCE_URLS:
            print(f"  [+] Scrapling loading {url} ...")
            session.fetch(url, page_action=grab)
            html = captured.get("html", "")
            if "theme_scriptspf" in html:
                break

    if "cf_clearance" not in captured.get("cookies", {}) or \
            "Performing security verification" in html or "theme_scriptspf" not in html:
        raise RuntimeError(
            "Cloudflare challenge was NOT cleared (still on the interstitial, or a stale "
            "cf_clearance in the profile). Retry with --headed; if it still fails add "
            "--fresh-profile. Headless mode is usually flagged."
        )
    return captured["cookies"], captured["ua"], html


class Client:
    """Holds cookies/UA/nonce; hands out per-thread curl_cffi sessions and
    re-bootstraps once (not once per worker) when Cloudflare returns 403."""

    def __init__(self, headless=True):
        self.headless = headless
        self.lock = threading.Lock()
        self.local = threading.local()
        self.generation = 0
        self.cookies, self.ua, self.nonce = {}, "", ""
        self.bootstrap()

    def bootstrap(self):
        cookies, ua, html = solve_with_scrapling(self.headless)
        m = NONCE_RE.search(html)
        if not m:
            raise RuntimeError("theme_scriptspf not found on bootstrap page.")
        nonce = json.loads(m.group(1)).get("pfget_listitems")
        if not nonce:
            raise RuntimeError("theme_scriptspf has no pfget_listitems key.")
        self.cookies, self.ua, self.nonce = cookies, ua, nonce
        self.generation += 1
        print(f"  [+] Bootstrapped (gen {self.generation}). Nonce: {nonce}")

    def rebootstrap(self, seen_generation):
        with self.lock:
            if self.generation == seen_generation:   # nobody else did it yet
                print("  [!] Re-bootstrapping Cloudflare session...")
                self.bootstrap()

    def session(self):
        s = getattr(self.local, "s", None)
        if s is None or getattr(self.local, "gen", None) != self.generation:
            s = cr.Session(impersonate="chrome")
            s.headers.update({"User-Agent": self.ua, "Accept-Language": "en-US,en;q=0.9"})
            for k, v in self.cookies.items():
                s.cookies.set(k, v, domain="www.qatarcid.com")
            self.local.s, self.local.gen = s, self.generation
        return s

    def request(self, method, url, **kw):
        """Retry wrapper. 429 honors Retry-After; 403 triggers one rebootstrap."""
        last = None
        rebooted = False
        attempt = 0
        while attempt < RETRY_LIMIT:
            attempt += 1
            gen = self.generation
            try:
                r = self.session().request(method, url, timeout=REQUEST_TIMEOUT, **kw)
                if r.status_code == 429:
                    try:
                        wait = float(r.headers.get("Retry-After", DEFAULT_RETRY_AFTER_SECONDS))
                    except ValueError:
                        wait = DEFAULT_RETRY_AFTER_SECONDS
                    print(f"  [!] 429 on {url} -- waiting {wait}s")
                    time.sleep(wait)
                    continue
                if r.status_code == 403:
                    if not rebooted:
                        rebooted = True
                        self.rebootstrap(gen)
                        attempt -= 1   # don't burn a retry on the rebootstrap
                        continue
                    raise RuntimeError(f"403 persists on {url}")
                r.raise_for_status()
                return r
            except Exception as e:
                last = e
                if attempt < RETRY_LIMIT:
                    time.sleep(RETRY_BACKOFF_SECONDS * attempt)
        raise last if last else RuntimeError("request failed")


# ---------------------------------------------------------------------------
# Site-specific requests
# ---------------------------------------------------------------------------
def fetch_search_page(client, page, page_size):
    payload = {
        "action": "pfget_listitems", "act": "search",
        "dt[jobskeyword]": "", "dt[field_company]": "", "dt[field_listingtype]": "",
        "dt[geolocation]": "", "dt[pointfinder_google_search_coord]": "",
        "dt[pointfinder_google_search_coord_unit]": "Mile", "dt[pointfinder_radius_search]": "",
        "dt[ne]": "", "dt[ne2]": "", "dt[sw]": "", "dt[sw2]": "",
        "dt[CR_NO]": "", "dt[QCCI_MEM_NO]": "", "dt[s]": "",
        "dt[serialized]": "1", "dt[action]": "pfs",
        "dtx[0][name]": "post_tags", "dtx[0][value]": "",
        "dtx[1][name]": "pointfinderltypes", "dtx[1][value]": "",
        "dtx[2][name]": "pointfinderlocations", "dtx[2][value]": "",
        "dtx[3][name]": "pointfinderconditions", "dtx[3][value]": "",
        "dtx[4][name]": "pointfinderitypes", "dtx[4][value]": "",
        "dtx[5][name]": "pointfinderfeatures", "dtx[5][value]": "",
        "ne": "", "sw": "", "ne2": "", "sw2": "", "cl": "", "grid": "",
        "pfg_orderby": "", "pfg_order": "",
        "pfg_number": "" if page_size == 14 else str(page_size),
        "pfcontainerdiv": ".pfsearchresults", "pfcontainershow": ".pfsearchgridview",
        "page": "" if page == 1 else str(page),
        "from": "halfmap", "security": client.nonce,
        "pflat": "undefined", "pflng": "undefined", "ohours": "",
    }
    headers = {
        "X-Requested-With": "XMLHttpRequest", "Origin": BASE,
        "Referer": f"{BASE}/", "Accept": "text/html, */*; q=0.01",
    }
    return client.request("POST", AJAX_URL, data=payload, headers=headers).text


def fetch_detail_page(client, url):
    return client.request("GET", url, headers={"Referer": f"{BASE}/"}).text


# ---------------------------------------------------------------------------
# Probe
# ---------------------------------------------------------------------------
def run_probe(client):
    html = fetch_search_page(client, 1, 14)
    print(f"  [probe] default: {len(parse_listing_cards(html))} cards, total={parse_found_posts(html)}")

    # Does the server honor a bigger pfg_number?
    for n in (75, 50):
        try:
            h = fetch_search_page(client, 1, n)
            got = len(parse_listing_cards(h))
            print(f"  [probe] pfg_number={n}: {got} cards", "<-- USE --page-size %d" % n if got == n else "(capped/ignored)")
            if got == n:
                break
        except Exception as e:
            print(f"  [probe] pfg_number={n} failed: {e}")

    # REST API check
    try:
        r = client.request("GET", f"{BASE}/wp-json/wp/v2/types")
        print(f"  [probe] wp-json types: {list(r.json().keys())}")
    except Exception as e:
        print(f"  [probe] wp-json/types unavailable: {e}")


# ---------------------------------------------------------------------------
# Sweeps
# ---------------------------------------------------------------------------
def run_listing_sweep(state, client, page_size):
    state.page_size = page_size
    first = fetch_search_page(client, 1, page_size)
    total = parse_found_posts(first)
    if total is None:
        raise RuntimeError("data-foundposts not found -- markup changed?")
    state.total_found = total
    state.total_pages = math.ceil(total / page_size)
    print(f"  [+] {total} listings / {state.total_pages} pages of {page_size}")

    for u, n in parse_listing_cards(first):
        state.listing_urls[u] = n
    state.pages_done.add(1)
    state.save(CHECKPOINT_PATH)

    for page in range(2, state.total_pages + 1):
        if page in state.pages_done:
            continue
        time.sleep(POLITE_DELAY_SECONDS)
        html = fetch_search_page(client, page, page_size)
        for u, n in parse_listing_cards(html):
            state.listing_urls[u] = n
        state.pages_done.add(page)
        if page % 5 == 0 or page == state.total_pages:
            state.save(CHECKPOINT_PATH)
            print(f"  [+] Page {page}/{state.total_pages} -- {len(state.listing_urls)} unique")

    state.save(CHECKPOINT_PATH)
    diff = abs(total - len(state.listing_urls)) / total * 100
    if diff > 2:
        print(f"  [!] WARNING: {len(state.listing_urls)} collected vs {total} reported ({diff:.1f}% off) -- check pagination.")
    else:
        print(f"  [+] Sanity check OK ({diff:.1f}% off).")


def run_detail_sweep(state, client, workers):
    if not state.listing_urls:
        print("  [!] No listing URLs in checkpoint; run the listing sweep first.")
        return
    stale = [u for u, r in state.records.items()
             if any("[email" in str(v) and "protected]" in str(v) for v in r.values())]
    for u in stale:
        state.records.pop(u, None)
        state.detail_done.discard(u)
    if stale:
        print(f"  [+] Re-fetching {len(stale)} records saved with obfuscated emails.")
    todo = [u for u in state.listing_urls if u not in state.detail_done]
    print(f"  [+] {len(todo)} detail pages to fetch ({len(state.detail_done)} done), workers={workers}")

    def work(url):
        time.sleep(POLITE_DELAY_SECONDS)
        return url, parse_detail_fields(fetch_detail_page(client, url))

    done = 0
    it = iter(todo)
    pending = set()
    ex = ThreadPoolExecutor(max_workers=workers)
    try:
        # Sliding window: never more than workers*2 tasks queued, so Ctrl+C is instant.
        while True:
            while len(pending) < workers * 2:
                try:
                    pending.add(ex.submit(work, next(it)))
                except StopIteration:
                    break
            if not pending:
                break
            finished = [f for f in list(pending) if f.done()]
            if not finished:
                time.sleep(0.2)
                continue
            for fut in finished:
                pending.discard(fut)
                try:
                    url, fields = fut.result()
                except Exception as e:
                    print(f"  [!] skipped one: {e}")
                    continue
                fields["_url"] = url
                fields["_search_card_name"] = state.listing_urls[url]
                state.records[url] = fields
                state.detail_done.add(url)
                done += 1
                if done % 25 == 0:
                    state.save(CHECKPOINT_PATH)
                    print(f"  [+] Detail {done}/{len(todo)} -- {fields['_search_card_name']}")
    except KeyboardInterrupt:
        print("\n  [!] Interrupted -- cancelling pending tasks and saving checkpoint...")
    finally:
        ex.shutdown(wait=False, cancel_futures=True)
        state.save(CHECKPOINT_PATH)
    print(f"  [+] Detail sweep stopped/complete: {len(state.records)} records saved.")


def export_to_excel(state, path=OUTPUT_XLSX):
    import openpyxl
    from openpyxl.utils import get_column_letter
    if not state.records:
        print("  [!] No records to export.")
        return
    preferred = ["_search_card_name", "Listing Type", "Location", "QCCI Membership Number",
                 "CR Number", "Company Type", "Address", "PO Box", "Phone", "Fax", "Email",
                 "Website", "Contact Person Mobile", "Contact Person", "Owner Name",
                 "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
    keys = set()
    for r in state.records.values():
        keys.update(r)
    ordered = [k for k in preferred if k in keys] + \
              sorted(k for k in keys if k not in preferred and k != "_url")
    if "_url" in keys:
        ordered.append("_url")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Qatar CID Companies"
    ws.append(ordered)
    for r in state.records.values():
        ws.append([r.get(k, "") for k in ordered])
    from openpyxl.styles import Font
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, k in enumerate(ordered, 1):
        m = max([len(k)] + [len(str(r.get(k, ""))) for r in state.records.values()])
        ws.column_dimensions[get_column_letter(i)].width = min(m + 2, 60)
    wb.save(path)
    print(f"  [+] Exported {len(state.records)} records to {path}")


def main():
    global CHECKPOINT_PATH
    ap = argparse.ArgumentParser()
    ap.add_argument("--probe", action="store_true")
    ap.add_argument("--listings-only", action="store_true")
    ap.add_argument("--details-only", action="store_true")
    ap.add_argument("--export-only", action="store_true")
    ap.add_argument("--headed", action="store_true", help="Show the browser for the CF solve")
    ap.add_argument("--fresh-profile", action="store_true", help="Delete ./scrapling_profile before solving")
    ap.add_argument("--page-size", type=int, default=14, help="pfg_number (try 75 after --probe)")
    ap.add_argument("--workers", type=int, default=1, help="parallel detail fetches (3-4 max)")
    ap.add_argument("--checkpoint", default=CHECKPOINT_PATH)
    ap.add_argument("--output", default=OUTPUT_XLSX)
    args = ap.parse_args()

    CHECKPOINT_PATH = args.checkpoint

    state = ScrapeState.load(args.checkpoint)
    print(f"Loaded checkpoint: {len(state.listing_urls)} URLs, {len(state.records)} records")

    if args.export_only:
        export_to_excel(state, args.output)
        return

    if args.fresh_profile and os.path.isdir(PROFILE_DIR):
        import shutil
        shutil.rmtree(PROFILE_DIR, ignore_errors=True)
        print("  [+] Removed stale browser profile.")
    if not args.headed:
        print("  [i] Running headless -- if the solve fails, rerun with --headed.")
    client = Client(headless=not args.headed)

    if args.probe:
        run_probe(client)
        return

    if not args.details_only:
        if state.pages_done and state.page_size != args.page_size:
            print(f"  [!] Checkpoint used page size {state.page_size}; resuming with that.")
            args.page_size = state.page_size
        run_listing_sweep(state, client, args.page_size)

    if not args.listings_only:
        run_detail_sweep(state, client, args.workers)

    export_to_excel(state, args.output)


if __name__ == "__main__":
    main()
